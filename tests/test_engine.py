import _support
import hashlib
import tempfile
import threading
import unittest
from pathlib import Path
from PIL import Image
import pymupdf
from pptx import Presentation
from engine import convert, inspect, Cancelled


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.image = self.root / '한글 & source.png'
        Image.new('RGBA', (41, 29), (20, 80, 180, 100)).save(self.image)

    def test_image_conversion_preserves_source(self):
        before = hashlib.sha256(self.image.read_bytes()).digest()
        result = convert(self.image, 'jpg', self.root)
        with Image.open(result) as image:
            self.assertEqual(image.size, (41, 29))
            self.assertEqual(image.mode, 'RGB')
        self.assertEqual(before, hashlib.sha256(self.image.read_bytes()).digest())

    def test_existing_output_is_not_overwritten(self):
        existing = self.image.with_suffix('.jpg')
        existing.write_bytes(b'keep this file')
        result = convert(self.image, 'jpg', self.root)
        self.assertNotEqual(result, existing)
        self.assertEqual(existing.read_bytes(), b'keep this file')

    def test_pdf_all_pages_and_slide_count(self):
        pdf = self.root / 'input.pdf'
        with pymupdf.open() as doc:
            for i in range(2):
                page = doc.new_page(width=160, height=120)
                page.insert_text((20, 50), f'Page {i + 1}')
            doc.save(pdf)
        pages = convert(pdf, 'png', self.root)
        self.assertEqual(len(list(pages.glob('*.png'))), 2)
        slides = convert(pdf, 'pptx', self.root)
        deck = Presentation(slides)
        self.assertEqual(len(deck.slides), 2)
        self.assertEqual(len(deck.slides[0].shapes), 1)

    def test_pdf_to_txt_extracts_pages_in_order(self):
        pdf = self.root / 'text.pdf'
        with pymupdf.open() as doc:
            for i in range(2):
                page = doc.new_page(width=160, height=120)
                page.insert_text((20, 50), f'Page {i + 1}')
            doc.save(pdf)
        self.assertIn('txt', inspect(pdf)['formats'])
        text = convert(pdf, 'txt', self.root).read_text(encoding='utf-8')
        self.assertLess(text.index('Page 1'), text.index('Page 2'))

    def test_cancel_leaves_no_partial_output(self):
        cancelled = threading.Event()
        cancelled.set()
        with self.assertRaises(Cancelled):
            convert(self.image, 'jpg', self.root, cancel=cancelled)
        self.assertFalse(self.image.with_suffix('.jpg').exists())
        self.assertFalse(list(self.root.glob('.converter-*')))

    def test_unsupported_output_and_corrupt_input(self):
        with self.assertRaises(ValueError):
            convert(self.image, 'mp3', self.root)
        corrupt = self.root / 'broken.png'
        corrupt.write_bytes(b'not an image')
        with self.assertRaises(Exception):
            inspect(corrupt)


if __name__ == '__main__':
    unittest.main()
