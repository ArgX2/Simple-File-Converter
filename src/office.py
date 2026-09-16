from pathlib import Path
import io
import os
import sys
import shutil
import subprocess
import time
import tempfile
import winreg
import pymupdf
from pptx import Presentation
from pptx.util import Inches

def backend():
    try:
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r'PowerPoint.Application\CLSID'):
            return 'PowerPoint', ''
    except OSError:
        pass
    candidates = [shutil.which('soffice')]
    for base in ['PROGRAMFILES', 'PROGRAMFILES(X86)']:
        candidates.append(str(Path(os.environ.get(base, 'C:/Program Files')) / 'LibreOffice/program/soffice.exe'))
    for candidate in candidates:
        if candidate and Path(candidate).is_file(): return 'LibreOffice', candidate
    return None, None

def office_convert(src, dst, cancel):
    from engine import check
    name, exe = backend()
    if not name: raise ValueError('PowerPoint 또는 LibreOffice를 설치한 뒤 다시 시도해 주세요.')
    with tempfile.TemporaryDirectory(prefix='sfc-office-') as temp:
        root = Path(temp)
        if name == 'PowerPoint':
            script = Path(getattr(sys, '_MEIPASS', Path(__file__).parent)) / 'source' / 'office.ps1'
            if not script.exists(): script = Path(__file__).parent / 'office.ps1'
            args = [str(Path(os.environ['SYSTEMROOT']) / 'System32/WindowsPowerShell/v1.0/powershell.exe'), '-NoProfile', '-NonInteractive', '-ExecutionPolicy', 'Bypass', '-File', str(script), '-InputFile', str(src.resolve()), '-OutputFile', str(dst.resolve()), '-Format', '32' if dst.suffix == '.pdf' else '1']
        else:
            fmt = 'pdf:impress_pdf_Export' if dst.suffix == '.pdf' else 'ppt:MS PowerPoint 97'
            args = [exe, '-env:UserInstallation=' + (root / 'profile').as_uri(), '--headless', '--convert-to', fmt, '--outdir', str(root), str(src.resolve())]
        with tempfile.TemporaryFile() as log:
            process = subprocess.Popen(args, stdout=log, stderr=log, creationflags=subprocess.CREATE_NO_WINDOW)
            start = time.monotonic()
            try:
                while process.poll() is None:
                    # Let the Office helper close its owned document before cancellation.
                    if time.monotonic() - start > 120:
                        raise ValueError('Office의 응답이 늦어지고 있습니다. 열린 대화상자를 확인한 뒤 다시 시도해 주세요.')
                    time.sleep(.1)
                if process.returncode:
                    log.seek(0)
                    raise ValueError('Office에서 변환하지 못했습니다. 앱에서 파일이 열리는지 확인해 주세요.\n\n상세 오류: ' + log.read().decode('utf-8', errors='replace')[-1200:])
                check(cancel)
                if name == 'LibreOffice':
                    result = root / (src.stem + dst.suffix)
                    if result.exists(): shutil.copy2(result, dst)
                if not dst.is_file() or not dst.stat().st_size:
                    raise ValueError('변환 결과를 만들지 못했습니다. 파일의 암호와 손상 여부, Office 상태를 확인해 주세요.')
            finally:
                if process.poll() is None: process.kill()
                process.wait()

def pdf_to_slides(src, dst, cancel, progress):
    from engine import check
    deck = Presentation()
    with pymupdf.open(src) as doc:
        if doc.needs_pass: raise ValueError('PDF의 암호를 해제한 뒤 다시 시도해 주세요.')
        if not len(doc): raise ValueError('PDF에 페이지가 없습니다. 원본 파일을 확인해 주세요.')
        first = doc[0].rect
        deck.slide_width = Inches(10)
        deck.slide_height = int(deck.slide_width * first.height / first.width)
        for i, page in enumerate(doc):
            check(cancel)
            progress(f'슬라이드 {i+1}/{len(doc)}')
            pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            scale = min(deck.slide_width/pix.width, deck.slide_height/pix.height)
            width, height = int(pix.width*scale), int(pix.height*scale)
            slide.shapes.add_picture(io.BytesIO(pix.tobytes('png')), (deck.slide_width-width)//2, (deck.slide_height-height)//2, width, height)
    if dst.suffix == '.pptx': deck.save(dst)
    else:
        intermediate = dst.with_suffix('.pptx')
        deck.save(intermediate)
        office_convert(intermediate, dst, cancel)
